"""用冻结正文、报告、题录和页级证据生成 Word；本模块不调用模型。"""

from __future__ import annotations

from collections.abc import Mapping, Sequence
from io import BytesIO
from typing import Any
from zipfile import ZipFile

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from apps.common.http import api_error
from apps.common.storage import StagedFile, discard_staged, stage_file
from apps.reviews.filenames import export_docx_filename

TEMPLATE_REVISION = "docx-template-v4-source-commentary"
DOCX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
MAX_DOCX_BYTES = 32 * 1024 * 1024
SUPPORT_LABELS = {
    "supported": "证据支持",
    "unsupported": "证据不支持",
    "uncertain": "无法确定",
    "not_applicable": "纯承接语",
}
VERIFICATION_LABELS = {
    "text_matched": "已匹配文本原文",
    "image_unverified": "图像转录未经人工核对",
    "human_verified": "人工核对",
}


def build_bibliography(run_snapshot: Mapping[str, Any]) -> list[dict[str, Any]]:
    """由正文首次引用的论文顺序编号；同一论文不同页只占一个文献号。"""
    papers = {p["paper_id"]: p for p in run_snapshot["papers"]}
    evidence = {e["evidence_id"]: e for e in run_snapshot["evidence"]}
    result: list[dict[str, Any]] = []
    seen: set[str] = set()
    for section in run_snapshot["sections"]:
        for block in section["blocks"]:
            for eid in block["evidence_ids"]:
                if eid not in evidence or evidence[eid]["paper_id"] not in papers:
                    raise api_error("INVALID_EVIDENCE", "引文没有对应的冻结题录")
                paper_id = evidence[eid]["paper_id"]
                if paper_id not in seen:
                    seen.add(paper_id)
                    result.append({"number": len(result) + 1, **papers[paper_id]})
    return result


def _style_document(doc: Any) -> None:
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.27), Inches(11.69)
    section.top_margin = section.bottom_margin = Inches(0.8)
    section.left_margin = section.right_margin = Inches(0.9)
    section.header_distance = section.footer_distance = Inches(0.35)
    for name in ("Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Header", "Footer"):
        style = doc.styles[name]
        style.font.name = "Times New Roman"
        style.font.color.rgb = RGBColor(0, 0, 0)
        fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
        for attribute in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
            fonts.attrib.pop(qn(f"w:{attribute}"), None)
        fonts.set(qn("w:eastAsia"), "宋体")
        for border in style.element.xpath("./w:pPr/w:pBdr"):
            border.getparent().remove(border)
    normal = doc.styles["Normal"]
    normal.font.size = Pt(11)
    normal.paragraph_format.line_spacing = 1.35
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.widow_control = True
    doc.styles["Title"].font.size = Pt(22)
    for name, size in (("Heading 1", 16), ("Heading 2", 12)):
        doc.styles[name].font.size = Pt(size)
        doc.styles[name].paragraph_format.keep_with_next = True
        doc.styles[name].paragraph_format.space_before = Pt(15)
        doc.styles[name].paragraph_format.space_after = Pt(7)
    doc.styles["Header"].font.size = Pt(9)
    section.header.paragraphs[0].text = ""
    footer = section.footer.paragraphs[0]
    footer.add_run("第 ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    footer.add_run(" 页")
    doc.core_properties.author = "LitReview Generator"
    doc.core_properties.last_modified_by = "LitReview Generator"


def render_docx(
    revision_snapshot: Mapping[str, Any],
    validation_report: Mapping[str, Any],
    bibliography: Sequence[Any],
    evidence: Sequence[Any],
) -> StagedFile:
    from apps.reviews.structure import (
        MISSING_CONTENT,
        commentary_notices,
        is_commentary_structure,
        section_heading,
    )

    if not is_commentary_structure(revision_snapshot["sections"]):
        raise api_error("STRUCTURAL_VALIDATION_FAILED", "Word 导出必须使用固定研究框架，请重新生成综述")
    doc = Document()
    _style_document(doc)
    doc.add_paragraph(revision_snapshot["title"], "Title")
    citations = revision_snapshot.get("source_citations", {})
    commentary = is_commentary_structure(revision_snapshot["sections"])
    notices = commentary_notices(revision_snapshot["sections"], validation_report.get("blocks"))
    for section in revision_snapshot["sections"]:
        key = section["section_key"]
        if key == "theoretical_basis":
            doc.add_heading("4 文献综述", level=1)
        if commentary and key == "achievements":
            doc.add_heading("5 文献述评", level=1)
        heading, level = section_heading(key, legacy=not commentary)
        doc.add_heading(heading, level=level)
        if commentary and key == "limitations":
            for notice in notices:
                doc.add_paragraph(notice)
        for block in section["blocks"]:
            paragraph = doc.add_paragraph(block["text"])
            sources = citations.get(block["block_id"])
            if block["text"] == MISSING_CONTENT and not block["evidence_ids"]:
                continue
            if not sources:
                raise api_error("INVALID_EVIDENCE", "正文缺少可靠的 PDF 页段行出处")
            labels = []
            for source in sources:
                lines = str(source["line_start"])
                if source["line_end"] != source["line_start"]:
                    lines += f"—{source['line_end']}"
                labels.append(
                    f"《{source['filename']}》第{source['page']}页，第{source['paragraph']}段，第{lines}行"
                )
            paragraph.add_run("（" + "；".join(labels) + "）")
    # Audit details remain in the database and document properties, outside
    # the article's fixed body structure.
    doc.core_properties.comments = (
        f"revision={revision_snapshot['revision_id']}; hash={revision_snapshot['content_hash']}; "
        f"validation={validation_report['id']}; result={validation_report['result']}; "
        f"template={TEMPLATE_REVISION}"
    )
    stream = BytesIO()
    doc.save(stream)
    stream.seek(0)
    staged = stage_file(
        stream,
        limit_bytes=MAX_DOCX_BYTES,
        original_filename=export_docx_filename(
            revision_snapshot.get("project_title", revision_snapshot["title"])
        ),
        media_type=DOCX_MEDIA_TYPE,
        suffix=".docx",
    )
    try:
        with ZipFile(staged.path) as archive:
            if archive.testzip() is not None:
                raise api_error("INTERNAL_ERROR", "Word 文件完整性检查失败")
        Document(staged.path)
    except BaseException:
        discard_staged(staged)
        raise
    return staged
