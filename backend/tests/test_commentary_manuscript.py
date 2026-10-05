"""Source modules stay literal; commentary is reasoned, cited and honestly marked."""

import copy
import hashlib
import json
import uuid

import pytest
from docx import Document
from tests import test_m10_reviews as m10
from tests import test_m11_reviews as m11

from apps.analysis.services import _run_scope
from apps.common.storage import open_asset
from apps.parsing.models import EvidenceChunk
from apps.reviews.grounding import ground_sections
from apps.reviews.models import ReviewRevision
from apps.reviews.schemas import sections_content_hash, validate_sections
from apps.reviews.structure import (
    LEGACY_SOURCE_KEYS,
    LEGACY_SOURCE_TITLES,
    SOURCE_KEYS,
    SOURCE_POLICY,
    commentary_notices,
)

pytestmark = pytest.mark.django_db(transaction=True)
run_ready = m10.run_ready
generated = m10.generated
validated = m11.validated


def save(client, g, sections):
    return client.post(
        f"/api/v1/reviews/{g.review.id}/revisions",
        {"base_revision_number": g.revision.revision_number, "sections": sections},
        format="json",
    )


def test_commentary_can_be_reasoned_instead_of_verbatim_and_has_real_source_locations(client_a, generated):
    sections = copy.deepcopy(generated.revision.sections)
    block = sections[-1]["blocks"][0]
    block["text"] = "建议比较不同样本中的适用边界，以验证前文结果能否推广；尚未完成这项研究。"
    reply = save(client_a, generated, sections)
    assert reply.status_code == 201, reply.json()
    revision = ReviewRevision.objects.get(pk=reply.json()["data"]["id"])
    source_citations = ground_sections(revision.run, revision.sections)
    assert len(block["evidence_ids"]) == 2
    assert {p["evidence_id"] for p in source_citations[block["block_id"]]} == set(block["evidence_ids"])
    assert all(
        p["page"] > 0 and p["paragraph"] > 0 and p["line_start"] > 0
        for p in source_citations[block["block_id"]]
    )


def test_commentary_rejects_unused_run_evidence(client_a, generated):
    sections = copy.deepcopy(generated.revision.sections)
    used = {
        eid
        for s in sections
        if s["section_key"] in SOURCE_KEYS
        for b in s["blocks"]
        for eid in b["evidence_ids"]
    }
    unused = (
        EvidenceChunk.objects.filter(
            parsed_document__in=generated.run.run_papers.values_list("parsed_document_id", flat=True)
        )
        .exclude(pk__in=used)
        .first()
    )
    assert unused is not None
    sections[-1]["blocks"][0]["evidence_ids"] = [str(unused.pk)]
    reply = save(client_a, generated, sections)
    assert reply.status_code == 422 and "前文" in reply.json()["error"]["message"]


@pytest.mark.parametrize("change", ["role", "number", "source_role", "transition_role", "prefix"])
def test_commentary_metadata_and_simulation_prefix_cannot_bypass_rules(client_a, generated, change):
    sections = copy.deepcopy(generated.revision.sections)
    if change == "role":
        sections[-1]["blocks"][0]["analysis_role"] = "achievement"
    elif change == "number":
        sections[-2]["blocks"][0]["controversy_number"] = True
    elif change == "source_role":
        sections[0]["blocks"][0]["analysis_role"] = "innovation"
    elif change == "transition_role":
        sections[-1]["blocks"] = [
            {
                **sections[-1]["blocks"][0],
                "kind": "transition",
                "text": "提供的素材中没有该内容。",
                "evidence_ids": [],
            }
        ]
    else:
        sections[-1]["blocks"][0]["text"] = "模拟草稿（需人工审阅）：" + sections[-1]["blocks"][0]["text"]
    assert save(client_a, generated, sections).status_code == 422


def test_second_controversy_needs_its_own_cause_and_metadata_is_hashed(generated):
    sections = copy.deepcopy(generated.revision.sections)
    assert commentary_notices(sections) == []
    original = sections_content_hash(sections)
    sections[-2]["blocks"][-1]["controversy_number"] = 1
    assert sections_content_hash(sections) != original
    assert any("争议证据不足" in s for s in commentary_notices(sections))
    assert validate_sections(sections, run_scope=_run_scope(generated.run)) == sections


def test_two_numbers_do_not_hide_a_controversy_that_failed_semantic_validation(
    client_a, generated, monkeypatch
):
    target = generated.revision.sections[-2]["blocks"][-1]["block_id"]

    def transform(rows):
        return [
            {
                **row,
                "support": "uncertain",
                "issues": ["样本差异不能在素材中确认，争议成因依据不足。"],
                "suggestion": "补充依据。",
            }
            if row["block_id"] == target
            else row
            for row in rows
        ]

    report, _ = m11.finish_validation(client_a, generated, monkeypatch, transform=transform)
    assert report.result == "needs_review"
    generated.report = report
    export, _ = m11.finish_export(client_a, generated, ack=True)
    with open_asset(export.file_asset) as stream:
        word = Document(stream)
    assert any("争议证据不足" in p.text for p in word.paragraphs)
    assert not any("模拟草稿" in p.text for p in word.paragraphs)


def test_complete_commentary_export_uses_fixed_hierarchy_without_insufficient_notice(client_a, validated):
    export, _ = m11.finish_export(client_a, validated)
    with open_asset(export.file_asset) as stream:
        word = Document(stream)
    assert [p.text for p in word.paragraphs if p.style.name.startswith("Heading")] == [
        "1 研究背景",
        "2 研究思路",
        "3 研究意义",
        "4 文献综述",
        "4.1 理论基础",
        "4.2 国内外研究现状",
        "5 文献述评",
        "5.1 已取得的成果",
        "5.2 文献的不足之处",
        "5.3 本文创新点",
    ]
    assert not any("争议证据不足" in p.text or "模拟草稿" in p.text for p in word.paragraphs)


def test_historical_six_slot_revision_remains_readable_and_hash_compatible(client_a, generated):
    by_key = {s["section_key"]: s for s in generated.revision.sections}
    sections = []
    for key in LEGACY_SOURCE_KEYS:
        section = copy.deepcopy(by_key[key if key != "achievements" else "background"])
        section["section_key"], section["title"] = key, LEGACY_SOURCE_TITLES[key]
        for block in section["blocks"]:
            block["block_id"] = str(uuid.uuid4())
        sections.append(section)
    expected_hash = hashlib.sha256(
        json.dumps(sections, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode()
    ).hexdigest()
    assert sections_content_hash(sections) == expected_hash
    generated.revision.sections = sections
    generated.revision.config_snapshot = {"source_policy": SOURCE_POLICY}
    generated.revision.content_hash = expected_hash
    generated.revision.save(update_fields=["sections", "config_snapshot", "content_hash"])
    assert validate_sections(sections, run_scope=_run_scope(generated.run)) == sections
    assert ground_sections(generated.run, sections)
    response = client_a.get(f"/api/v1/reviews/{generated.review.id}/revisions/{generated.revision.id}")
    assert response.status_code == 200
    reply = save(client_a, generated, sections)
    assert reply.status_code == 201, reply.json()


def test_historical_six_slot_manuscript_cannot_create_a_new_export(client_a, generated, monkeypatch):
    by_key = {s["section_key"]: s for s in generated.revision.sections}
    sections = []
    for key in LEGACY_SOURCE_KEYS:
        section = copy.deepcopy(by_key[key if key != "achievements" else "background"])
        section["section_key"], section["title"] = key, LEGACY_SOURCE_TITLES[key]
        for block in section["blocks"]:
            block["block_id"] = str(uuid.uuid4())
        sections.append(section)
    generated.revision.sections = sections
    generated.revision.config_snapshot = {"source_policy": SOURCE_POLICY}
    generated.revision.content_hash = sections_content_hash(sections)
    generated.revision.save(update_fields=["sections", "config_snapshot", "content_hash"])
    generated.report, _ = m11.finish_validation(client_a, generated, monkeypatch)
    reply = m11.accept_export(client_a, generated)
    assert reply.status_code == 422 and "旧结构" in reply.json()["error"]["message"]
