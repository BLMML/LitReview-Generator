"""Fixed framework, unchanged source wording and PDF paragraph/line citations."""

import copy
import uuid
from types import SimpleNamespace

import pytest
from docx import Document
from tests import test_m10_reviews as m10
from tests import test_m11_reviews as m11

from apps.common.storage import open_asset
from apps.parsing import locations
from apps.reviews.grounding import ground_sections
from apps.reviews.models import ReviewRevision
from apps.reviews.structure import MISSING_CONTENT, SECTION_KEYS, SECTION_TITLES

pytestmark = pytest.mark.django_db(transaction=True)
run_ready = m10.run_ready
generated = m10.generated
validated = m11.validated


@pytest.mark.parametrize("quote", ["improve outcomes.", "5 points.", "It does not improve"])
def test_partial_sentence_cannot_drop_negation_or_condition(monkeypatch, quote):
    monkeypatch.setattr(
        locations,
        "page_lines",
        lambda *args: (
            {"text": "It does not improve outcomes. It increased by 3.5 points.", "paragraph": 1, "line": 1},
        ),
    )
    with pytest.raises(ValueError, match="完整句子"):
        locations.locate_quote(SimpleNamespace(storage_key="x", sha256="x"), 1, quote)


def test_canonical_word_boundaries_and_multiline_locations(monkeypatch):
    monkeypatch.setattr(
        locations,
        "page_lines",
        lambda *args: (
            {"text": "The result does not", "paragraph": 2, "line": 3},
            {"text": "support that claim.", "paragraph": 2, "line": 4},
            {"text": "The sample is small.", "paragraph": 3, "line": 1},
        ),
    )
    source, spans = locations.locate_quote(
        SimpleNamespace(storage_key="x", sha256="x"),
        4,
        "Theresultdoesnotsupportthatclaim.The sample is small.",
    )
    assert source == "The result does not support that claim. The sample is small."
    assert spans == [
        {"page": 4, "paragraph": 2, "line_start": 3, "line_end": 4},
        {"page": 4, "paragraph": 3, "line_start": 1, "line_end": 1},
    ]


@pytest.mark.parametrize(
    "change",
    ["subjective", "mock_prefix", "fake_location", "crop", "title", "order", "extra_source", "transition"],
)
def test_save_rejects_non_source_content_or_framework_change(client_a, generated, change):
    sections = copy.deepcopy(generated.revision.sections)
    block = sections[0]["blocks"][0]
    if change == "subjective":
        block["text"] += " This is a pioneering achievement."
    elif change == "mock_prefix":
        block["text"] = "模拟草稿（需人工审阅）：" + block["text"]
    elif change == "fake_location":
        block["text"] += "（第999页，第999段，第999行）"
    elif change == "crop":
        block["text"] = "improves sequence modelling on small corpora."
    elif change == "title":
        sections[0]["title"] = "自由标题"
    elif change == "order":
        sections[0], sections[1] = sections[1], sections[0]
    elif change == "extra_source":
        block["evidence_ids"] += sections[0]["blocks"][1]["evidence_ids"]
    else:
        sections[0]["blocks"] = [
            {**block, "kind": "transition", "text": "以下进行总结。", "evidence_ids": []}
        ]
    reply = client_a.post(
        f"/api/v1/reviews/{generated.review.id}/revisions",
        {"base_revision_number": 1, "sections": sections},
        format="json",
    )
    assert reply.status_code == 422
    assert reply.json()["error"]["code"] == "STRUCTURAL_VALIDATION_FAILED"
    assert generated.review.revisions.count() == 1


def test_missing_modules_export_exact_notice_and_fixed_hierarchy(client_a, generated, monkeypatch):
    sections = [
        {
            "section_key": key,
            "title": SECTION_TITLES[key],
            "blocks": [
                {
                    "block_id": str(uuid.uuid4()),
                    "kind": "transition",
                    "text": MISSING_CONTENT,
                    "evidence_ids": [],
                }
            ],
        }
        for key in SECTION_KEYS
    ]
    reply = client_a.post(
        f"/api/v1/reviews/{generated.review.id}/revisions",
        {"base_revision_number": 1, "sections": sections},
        format="json",
    )
    assert reply.status_code == 201, reply.json()
    generated.revision = ReviewRevision.objects.get(pk=reply.json()["data"]["id"])
    calls = []
    generated.report, _ = m11.finish_validation(
        client_a, generated, monkeypatch, calls=calls, support="not_applicable"
    )
    assert generated.report.result == "passed"
    export, _ = m11.finish_export(client_a, generated)
    with open_asset(export.file_asset) as stream:
        word = Document(stream)
    assert [(p.style.name, p.text) for p in word.paragraphs if p.style.name.startswith("Heading")] == [
        ("Heading 1", "1 研究背景"),
        ("Heading 1", "2 研究意义"),
        ("Heading 1", "3 研究思路"),
        ("Heading 1", "4 文献综述"),
        ("Heading 2", "4.1 理论基础"),
        ("Heading 2", "4.2 国内外研究现状"),
        ("Heading 1", "5 已取得研究成果"),
    ]
    assert [p.text for p in word.paragraphs].count(MISSING_CONTENT) == 6
    assert not any("模拟草稿" in p.text or "参考文献" in p.text for p in word.paragraphs)


def test_word_body_only_contains_source_text_and_real_pdf_locations(client_a, validated):
    citations = ground_sections(validated.run, validated.revision.sections)
    assert citations == validated.revision.config_snapshot["source_citations"]
    export, _ = m11.finish_export(client_a, validated)
    with open_asset(export.file_asset) as stream:
        word = Document(stream)
    body = [p.text for p in word.paragraphs if p.style.name == "Normal" and p.text]
    blocks = [b for s in validated.revision.sections for b in s["blocks"]]
    assert len(body) == len(blocks)
    for text, block in zip(body, blocks, strict=True):
        assert text.startswith(block["text"] + "（《")
        for position in citations[block["block_id"]]:
            assert position["filename"] in text
            assert position["page"] > 0 and position["paragraph"] > 0
            assert position["line_end"] >= position["line_start"] > 0
            assert f"第{position['page']}页，第{position['paragraph']}段，第" in text
    assert "passed" in word.core_properties.comments


def test_semantic_validation_receives_fixed_module_context(client_a, generated, monkeypatch):
    calls = []
    m11.finish_validation(client_a, generated, monkeypatch, calls=calls)
    expected = {
        b["block_id"]: (s["section_key"], s["title"])
        for s in generated.revision.sections
        for b in s["blocks"]
    }
    assert calls
    for request in calls:
        for block in request.payload["blocks_with_evidence"]:
            assert (block["section_key"], block["section_title"]) == expected[block["block_id"]]
            assert block["source_policy"] == "pdf-source-only-v1"


def test_old_manuscript_requires_regeneration_before_new_export(client_a, generated, monkeypatch):
    m11.use_legacy_revision(generated)
    generated.report, _ = m11.finish_validation(client_a, generated, monkeypatch)
    reply = m11.accept_export(client_a, generated)
    assert reply.status_code == 422 and "旧结构" in reply.json()["error"]["message"]


def test_quotable_passages_exclude_cut_clauses_at_chunk_edges(monkeypatch):
    monkeypatch.setattr(
        locations,
        "page_lines",
        lambda *args: (
            {
                "text": (
                    "The original first sentence is complete. "
                    "It does not improve outcomes. A cut final clause"
                ),
                "paragraph": 1,
                "line": 1,
            },
        ),
    )
    asset = SimpleNamespace(storage_key="x", sha256="x")
    passages = locations.quotable_passages(
        asset, 1, "complete. It does not improve outcomes. A cut final clause"
    )
    assert passages == ["It does not improve outcomes."]
    passages = locations.quotable_passages(
        asset, 1, "The original first sentence is complete. It does not improve outcomes. A cut final clause"
    )
    assert set(passages) == {
        "The original first sentence is complete.",
        "It does not improve outcomes.",
        "The original first sentence is complete. It does not improve outcomes.",
    }
    assert not any("cut final" in passage for passage in passages)


def test_quotable_passages_never_shorten_an_oversized_single_sentence(monkeypatch):
    text = "It does not improve outcomes when " + "conditions " * 80 + "remain unchanged."
    monkeypatch.setattr(locations, "page_lines", lambda *args: ({"text": text, "paragraph": 1, "line": 1},))
    assert locations.quotable_passages(SimpleNamespace(storage_key="x", sha256="x"), 1, text) == []


@pytest.mark.parametrize("indented", [False, True])
def test_page_start_continuation_is_not_offered_as_an_intact_sentence(monkeypatch, indented):
    text = "a clause continuing the previous page. This complete sentence preserves its meaning."
    monkeypatch.setattr(
        locations,
        "page_lines",
        lambda *args: (
            {"text": "Journal header", "paragraph": 1, "line": 1},
            {"text": text, "paragraph": 2, "line": 1, "indented": indented},
        ),
    )
    passages = locations.quotable_passages(SimpleNamespace(storage_key="x", sha256="x"), 2, text)
    assert "This complete sentence preserves its meaning." in passages
    assert ("a clause continuing the previous page." in passages) is indented


@pytest.mark.parametrize(
    "source",
    [
        "摘要：The whole sentence preserves its meaning.",
        "Abstract: The whole sentence preserves its meaning.",
        "The preceding sentence ends.①The whole sentence preserves its meaning.",
    ],
)
def test_abstract_label_and_preceding_footnote_do_not_hide_sentence_boundary(monkeypatch, source):
    monkeypatch.setattr(
        locations,
        "page_lines",
        lambda *args: ({"text": source, "paragraph": 1, "line": 1},),
    )
    quote = "The whole sentence preserves its meaning."
    original, spans = locations.locate_quote(SimpleNamespace(storage_key="x", sha256="x"), 1, quote)
    assert original == quote and len(spans) == 1


def test_ordinary_colon_does_not_allow_dropping_a_condition(monkeypatch):
    monkeypatch.setattr(
        locations,
        "page_lines",
        lambda *args: ({"text": "Under these conditions: it works.", "paragraph": 1, "line": 1},),
    )
    with pytest.raises(ValueError, match="完整句子"):
        locations.locate_quote(SimpleNamespace(storage_key="x", sha256="x"), 1, "it works.")
