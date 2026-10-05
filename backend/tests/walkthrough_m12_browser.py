"""Explicit M12 browser integration run (not collected by default).

Real Django HTTP/CSRF/session, PostgreSQL, parsing, index, services, gateway,
ledger and DOCX; the default mock provider is exercised without replacement; only queue delivery
is replaced by a local execution thread.
No resource status is manually promoted. Run with pytest on this exact file.
"""

from __future__ import annotations

import json
import os
import shutil
import subprocess
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path

import pytest
from django.contrib.auth import get_user_model
from django.db import close_old_connections, connections
from docx import Document

from apps.jobs.models import Job
from apps.jobs.tasks import execute_job
from apps.literature.models import Paper
from apps.projects.models import Project
from apps.reviews.models import Export, ReviewRevision

pytestmark = pytest.mark.django_db(transaction=True)


def test_full_browser_flow(live_server, settings, monkeypatch, tmp_path):
    settings.ALLOWED_HOSTS = ["localhost", "127.0.0.1", "testserver"]
    settings.CSRF_TRUSTED_ORIGINS = ["http://127.0.0.1:5174"]
    settings.SESSION_COOKIE_SECURE = False
    settings.CSRF_COOKIE_SECURE = False
    # Use the deployment input limit; the browser flow must not rely on a
    # test-only inflated token bound to process these two real papers.
    get_user_model().objects.create_user("m12-browser", password="m12-test-only")
    futures = []

    def execute(job_id):
        close_old_connections()
        try:
            execute_job(str(job_id))
        finally:
            connections.close_all()

    root = Path(settings.REPO_ROOT)
    papers_dir = root / ".data/m12-qa/papers"
    supplied = json.loads(os.environ.get("M12_PDF_PATHS", "[]"))
    input_files = supplied or [str(papers_dir / name) for name in ("attention.pdf", "bert.pdf")]
    assert all(Path(name).is_file() for name in input_files), (
        "See M12 acceptance doc: download the two public test PDFs first."
    )
    with ThreadPoolExecutor(max_workers=1) as pool:
        monkeypatch.setattr(
            "apps.jobs.services.dispatch", lambda job: futures.append(pool.submit(execute, job.id))
        )
        env = {
            **os.environ,
            "M12_LIVE_TEST": "1",
            "M12_API_PROXY_TARGET": live_server.url,
            "M12_PAPERS_DIR": str(papers_dir),
            "M12_INPUT_FILES": json.dumps(input_files, ensure_ascii=False),
            "M12_PROJECT_TITLE": os.environ.get("M12_PROJECT_TITLE", "文献全流程浏览器联调"),
            "M12_RESEARCH_QUESTION": os.environ.get(
                "M12_RESEARCH_QUESTION", "社会工作与志愿服务在基层治理中如何协同？"
            ),
            "M12_REVIEW_TITLE": os.environ.get("M12_REVIEW_TITLE", "社会治理文献全流程测试"),
            "M12_QA_DIR": os.environ.get("M12_QA_DIR", str(root / ".data/m12-qa")),
            "LLM_MODE": "mock",
        }
        result = subprocess.run(  # noqa: S603 - executable resolved from trusted local Node installation
            [shutil.which("npm.cmd" if os.name == "nt" else "npm") or "npm", "run", "test:e2e"],
            cwd=root / "frontend",
            env=env,
            capture_output=True,
            text=True,
            encoding="utf-8",
            errors="replace",
            timeout=300,
        )
        (root / ".data/m12-qa/browser-live.txt").write_text(result.stdout + result.stderr, encoding="utf-8")
        for future in futures:
            future.result(timeout=30)
        assert result.returncode == 0, (
            result.stdout + result.stderr + str(list(Job.objects.values("kind", "status", "error")))
        )
    assert not Job.objects.exclude(status="succeeded").exists(), list(
        Job.objects.values("kind", "status", "error")
    )
    analysis_job = Job.objects.get(kind="analyze")
    assert analysis_job.stage == "finalize"
    assert analysis_job.progress == {"completed": 2, "total": 2, "unit": "paper"}
    for kind, stage, total in (
        ("retrieve", "publish", 4), ("generate_review", "commit_revision", 6),
        ("validate_review", "summarize", 3), ("export_docx", "publish", 2),
    ):
        job = Job.objects.get(kind=kind)
        assert job.stage == stage
        assert job.progress["completed"] == job.progress["total"] == total
    assert ReviewRevision.objects.count() == 2
    assert Paper.objects.count() == Paper.objects.filter(deleted_at__isnull=True).count() == 2
    assert Paper.objects.filter(version=3, included=True).count() == 1
    assert Project.objects.count() == Project.objects.filter(deleted_at__isnull=False).count() == 1
    from apps.llm.models import ModelCall

    assert all(
        call.model_snapshot.get("default_mock_version") == "request-bound-v2-source-only"
        for call in ModelCall.objects.all()
    )
    assert not ModelCall.objects.filter(outcome="failed").exists()
    exported = Export.objects.select_related("file_asset", "revision").get(status="ready")
    assert exported.revision.revision_number == 2 and exported.warnings
    from apps.common.storage import open_asset

    with open_asset(exported.file_asset) as stream:
        word = Document(stream)
    text = "\n".join(p.text for p in word.paragraphs)
    assert "模拟草稿（需人工审阅）：" not in text and "人工补充" not in text
    assert "needs_review" not in text
    assert not any("参考文献" in p.text for p in word.paragraphs)
    assert "4 文献综述" in text and "4.1 理论基础" in text and "4.2 国内外研究现状" in text
