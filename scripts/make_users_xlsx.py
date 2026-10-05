from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

rows = [
    ("alexchen",  "Kp7Rm2Xq"),
    ("brianlee",  "Tv4Nc9Zs"),
    ("carolwang", "Hd6Gy3Bw"),
    ("davidliu",  "Mq8Ft5Jv"),
    ("emmasun",   "Rc2Wp7Lk"),
]

wb = Workbook()
ws = wb.active
ws.title = "Users"

headers = ["No.", "Username", "Password"]
ws.append(headers)
for i, (user, pwd) in enumerate(rows, start=1):
    ws.append([i, user, pwd])

header_fill = PatternFill("solid", fgColor="1F4E79")
header_font = Font(bold=True, color="FFFFFF", size=11)
thin = Side(style="thin", color="B7C4D6")
border = Border(left=thin, right=thin, top=thin, bottom=thin)

for cell in ws[1]:
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.border = border

for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=3):
    for cell in row:
        cell.border = border
        cell.alignment = Alignment(
            horizontal="center" if cell.column in (1,) else "left",
            vertical="center",
        )
    row[1].font = Font(name="Consolas", size=11)
    row[2].font = Font(name="Consolas", size=11)

for col, width in zip("ABC", (6, 18, 18)):
    ws.column_dimensions[col].width = width
ws.row_dimensions[1].height = 22
ws.freeze_panes = "A2"

wb.save(r"D:\LitReview Generator\user_accounts.xlsx")
print("saved")
